import csv
import io
import zipfile
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

from backend.chat import export_file_tool
from backend.chat.export_file_tool import TOOL_NAME, build_file, safe_filename
from backend.chat.local_tools import LocalTools, merge_tool_lists

COLUMNS = ["ticker", "exchange", "held", "last_date", "days_stale"]
ROWS = [["VOD.L", "L", True, "2026-09-01", 33], ["AAPL", "N", False, None, 120]]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _args(fmt, **extra):
    return {"filename": "stale-audit", "format": fmt, "columns": COLUMNS, "rows": ROWS, **extra}


def test_csv_has_a_header_row_and_one_line_per_row():
    exported = build_file(_args("csv"))
    assert exported.filename == "stale-audit.csv"
    assert exported.media_type == "text/csv"
    lines = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert lines == [COLUMNS, ["VOD.L", "L", "True", "2026-09-01", "33"], ["AAPL", "N", "False", "", "120"]]


def test_csv_neutralises_formulas_but_not_negative_numbers():
    exported = build_file(
        {"filename": "x", "format": "csv", "columns": ["a", "b", "c"], "rows": [["=HYPERLINK(1)", "-3.5", -2]]}
    )
    _, row = list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))
    assert row == ["'=HYPERLINK(1)", "-3.5", "-2"]


def test_xlsx_keeps_types_and_stores_formula_text_as_text():
    rows = ROWS + [["=1+1", "L", False, None, 0]]
    exported = build_file(_args("xlsx", rows=rows, title="Stale data: audit"))
    assert exported.filename == "stale-audit.xlsx"
    sheet = load_workbook(io.BytesIO(exported.content)).active
    assert sheet.title == "Stale data  audit"
    values = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert values[0] == COLUMNS
    assert values[1] == ["VOD.L", "L", True, "2026-09-01", 33]
    assert sheet.cell(row=4, column=1).data_type == "s"
    assert values[3][0] == "=1+1"


def test_docx_is_a_word_package_with_title_and_table():
    exported = build_file(_args("docx", title="Stale <data> & audit"))
    assert exported.filename == "stale-audit.docx"
    with zipfile.ZipFile(io.BytesIO(exported.content)) as package:
        assert {"[Content_Types].xml", "_rels/.rels", "word/document.xml"} <= set(package.namelist())
        document = ElementTree.fromstring(package.read("word/document.xml"))
    paragraphs = document.find(f"{W}body").findall(f"{W}p")
    assert "".join(t.text for t in paragraphs[0].iter(f"{W}t")) == "Stale <data> & audit"
    rows = document.find(f"{W}body").find(f"{W}tbl").findall(f"{W}tr")
    cells = [["".join(t.text or "" for t in tc.iter(f"{W}t")) for tc in tr.findall(f"{W}tc")] for tr in rows]
    assert cells == [COLUMNS, ["VOD.L", "L", "True", "2026-09-01", "33"], ["AAPL", "N", "False", "", "120"]]


def test_control_characters_do_not_break_xml_formats():
    args = {"filename": "x", "columns": ["a\x01"], "rows": [["bad\x07value"]]}
    build_file({**args, "format": "xlsx"})
    with zipfile.ZipFile(io.BytesIO(build_file({**args, "format": "docx"}).content)) as package:
        ElementTree.fromstring(package.read("word/document.xml"))


@pytest.mark.parametrize(
    "name, expected",
    [
        ("report.csv", "report"),
        ("../../etc/passwd", "passwd"),
        ("C:\\temp\\audit.XLSX", "audit"),
        ("stale data / 2026", "2026"),
        ("<script>", "script"),
        ("", "export"),
        (None, "export"),
    ],
)
def test_safe_filename(name, expected):
    assert safe_filename(name) == expected


@pytest.mark.parametrize(
    "args, message",
    [
        ({"format": "pdf", "columns": ["a"], "rows": []}, "format must be"),
        ({"format": "csv", "columns": [], "rows": []}, "columns must be"),
        ({"format": "csv", "columns": ["a", "b"], "rows": [["only one"]]}, "row 1 must be"),
        ({"format": "csv", "columns": ["a"], "rows": [[{"x": 1}]]}, "nested value"),
        ({"format": "csv", "columns": ["a"], "rows": "nope"}, "rows must be"),
    ],
)
def test_bad_arguments_are_reported_to_the_model(args, message):
    text, is_error, exported = export_file_tool.call({"filename": "x", **args})
    assert is_error
    assert message in text
    assert exported is None


def test_row_and_size_limits(monkeypatch):
    too_many = {"filename": "x", "format": "csv", "columns": ["a"], "rows": [[1]] * (export_file_tool.MAX_ROWS + 1)}
    assert "at most" in export_file_tool.call(too_many)[0]
    monkeypatch.setattr(export_file_tool, "MAX_FILE_BYTES", 10)
    text, is_error, _ = export_file_tool.call(_args("csv"))
    assert is_error and "the limit is 10" in text


def test_local_tools_offer_export_only_when_enabled():
    assert TOOL_NAME not in [tool.name for tool in LocalTools().tools()]
    assert not LocalTools().handles(TOOL_NAME)
    local = LocalTools(file_exports=True)
    assert TOOL_NAME in [tool.name for tool in merge_tool_lists([], local)]
    assert local.handles(TOOL_NAME)


def test_local_tools_collect_exported_files():
    local = LocalTools(file_exports=True)
    text, is_error = local.call(TOOL_NAME, _args("csv"))
    assert not is_error
    assert "stale-audit.csv" in text and "download link" in text
    assert [item.filename for item in local.files] == ["stale-audit.csv"]


def test_local_tools_limit_files_per_turn(monkeypatch):
    monkeypatch.setattr(export_file_tool, "MAX_FILES_PER_TURN", 1)
    local = LocalTools(file_exports=True)
    assert not local.call(TOOL_NAME, _args("csv"))[1]
    text, is_error = local.call(TOOL_NAME, _args("xlsx"))
    assert is_error and "at most 1 files" in text
    assert len(local.files) == 1


def test_export_tool_can_be_switched_off(monkeypatch):
    from backend import config_module

    monkeypatch.setattr(config_module.config, "mcp_tools", {TOOL_NAME: False})
    assert TOOL_NAME not in [tool.name for tool in merge_tool_lists([], LocalTools(file_exports=True))]
